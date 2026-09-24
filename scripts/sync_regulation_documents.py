#!/usr/bin/env python3
"""Download official regulation attachments and build searchable Markdown.

The XLSX files are inventories only.  This script materializes the attachment
text so the runtime can index the actual document body instead of merely
knowing that a title and URL exist.
"""

from __future__ import annotations

import argparse
import html
import re
import subprocess
import sys
import tempfile
import urllib.request
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "crawler_data"
ADMIN_XLSX = DATA / "北大行政單位法規彙整.xlsx"
ACADEMIC_XLSX = DATA / "北大學術單位法規彙整.xlsx"
FAILED_TEXT_MARKERS = ("未能抽出全文", "無法擷取內文", "擷取失敗")
MIN_SEARCHABLE_CHARS = 200
LINK_OVERRIDES = {
    # The XLSX attachment points to a retired DGBAS listing that now returns a
    # Cloudflare challenge.  Keep the current official law-content endpoint.
    "https://www.dgbas.gov.tw/lp.asp?ctNode=4938&CtUnit=1412&BaseDSD=7":
        "https://law.dgbas.gov.tw/LawContent.aspx?id=FL036372",
    "https://www.dgbas.gov.tw/lp.asp?ctNode=6337&CtUnit=2246&BaseDSD=7":
        "https://law.dgbas.gov.tw/LawContent.aspx?id=FL017693",
}


@dataclass(frozen=True)
class InventoryItem:
    title: str
    url: str
    tags: str = ""
    updated_date: str = ""


TARGETS = {
    "hr": ("人事室", DATA / "hr_regulations.md"),
    "oga": ("總務處", DATA / "oga_regulations.md"),
}


def normalize_title(value: str) -> str:
    value = re.sub(r"\.(pdf|docx?|odt|ods)$", "", str(value or "").strip(), flags=re.I)
    return re.sub(r"\s+", "", value)


def load_admin_items(dept_label: str) -> list[InventoryItem]:
    workbook = openpyxl.load_workbook(ADMIN_XLSX, read_only=True, data_only=True)
    try:
        rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
        headers = [str(value or "") for value in next(rows)]
        items = []
        for row in rows:
            record = dict(zip(headers, row))
            if str(record.get("處室") or "").strip() != dept_label:
                continue
            title = str(record.get("法規名稱") or "").strip()
            url = str(record.get("檔案連結") or record.get("網址") or "").strip()
            if title and url:
                items.append(InventoryItem(
                    title=title,
                    url=url,
                    tags=str(record.get("標籤") or "").strip(),
                    updated_date=str(record.get("上傳日期") or "").strip(),
                ))
        return items
    finally:
        workbook.close()


def load_missing_ge_items() -> list[InventoryItem]:
    source_text = (DATA / "cge_content.md").read_text(encoding="utf-8-sig")
    existing_titles = {
        normalize_title(match.group(1))
        for match in re.finditer(r"^##\s+(.+?)\s*$", source_text, re.M)
    }
    workbook = openpyxl.load_workbook(ACADEMIC_XLSX, read_only=True, data_only=True)
    try:
        rows = workbook["通識教育中心"].iter_rows(values_only=True)
        headers = [str(value or "") for value in next(rows)]
        items = []
        for row in rows:
            record = dict(zip(headers, row))
            title = str(record.get("title") or "").strip()
            if not title or normalize_title(title) in existing_titles:
                continue
            url = str(record.get("file_url") or record.get("source_page") or "").strip()
            if url:
                items.append(InventoryItem(
                    title=title,
                    url=url,
                    tags=str(record.get("tags") or "").strip(),
                    updated_date=str(record.get("updated_date") or "").strip(),
                ))
        return items
    finally:
        workbook.close()


def download(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": "NTPU-AIA-Knowledge-Sync/1.0"})
    with urllib.request.urlopen(request, timeout=90) as response:
        content = response.read()
    if len(content) < 32:
        raise ValueError(f"downloaded file is unexpectedly small ({len(content)} bytes)")
    return content


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" +\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(content: bytes) -> list[tuple[str, str]]:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as handle:
        handle.write(content)
        handle.flush()
        reader = PdfReader(handle.name)
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = _clean_text(page.extract_text() or "")
            if text:
                pages.append((f"Page {number}", text))
        page_count = len(reader.pages)
    searchable_text = re.sub(r"\s+", "", "\n".join(text for _, text in pages))
    links = extract_pdf_links(content)
    # Several inventory attachments are one-page Word exports containing only
    # a linked title.  The actual regulation lives at the embedded government
    # URL, so a successful PDF text extraction is not sufficient by itself.
    should_follow_link = page_count == 1 and len(searchable_text) < 200 and bool(links)
    if not should_follow_link and len(searchable_text) >= 20:
        return pages
    for linked_url in links:
        linked_url = LINK_OVERRIDES.get(linked_url, linked_url)
        try:
            linked_pages = extract_attachment(download(linked_url), linked_url, follow_pdf_links=False)
            if linked_pages:
                return [("官方連結", f"延伸官方來源：{linked_url}"), *linked_pages]
        except Exception:
            continue
    return extract_pdf_with_ocr(content)


def extract_pdf_links(content: bytes) -> list[str]:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as handle:
        handle.write(content)
        handle.flush()
        reader = PdfReader(handle.name)
        links = []
        for page in reader.pages:
            for annotation in page.get("/Annots") or []:
                action = annotation.get_object().get("/A") or {}
                url = str(action.get("/URI") or "").strip()
                if url.startswith(("https://", "http://")) and url not in links:
                    links.append(url)
        return links


def extract_pdf_with_ocr(content: bytes) -> list[tuple[str, str]]:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise ValueError(
            "PDF is image-only; install rapidocr_onnxruntime to enable OCR"
        ) from exc

    with tempfile.TemporaryDirectory(prefix="ntpu-reg-ocr-") as tmp_dir:
        tmp = Path(tmp_dir)
        pdf_path = tmp / "source.pdf"
        pdf_path.write_bytes(content)
        output_prefix = tmp / "page"
        subprocess.run(
            ["pdftoppm", "-jpeg", "-r", "180", str(pdf_path), str(output_prefix)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
        )
        image_paths = sorted(tmp.glob("page-*.jpg"))
        if not image_paths:
            raise ValueError("OCR renderer produced no page images")
        engine = RapidOCR()
        pages = []
        for number, image_path in enumerate(image_paths, start=1):
            result, _ = engine(str(image_path))
            lines = [str(row[1]).strip() for row in (result or []) if len(row) > 1 and str(row[1]).strip()]
            text = _clean_text("\n".join(lines))
            if text:
                pages.append((f"Page {number} (OCR)", text))
        if len(re.sub(r"\s+", "", "\n".join(text for _, text in pages))) < MIN_SEARCHABLE_CHARS:
            raise ValueError("OCR produced no usable searchable text")
        return pages


def extract_open_document(content: bytes) -> list[tuple[str, str]]:
    with tempfile.NamedTemporaryFile(suffix=".zip") as handle:
        handle.write(content)
        handle.flush()
        with zipfile.ZipFile(handle.name) as archive:
            xml_bytes = archive.read("content.xml")
    root = ElementTree.fromstring(xml_bytes)
    chunks = []
    for element in root.iter():
        text = "".join(element.itertext()).strip()
        if text and (element.tag.endswith("}p") or element.tag.endswith("}h")):
            chunks.append(text)
    return [("全文", _clean_text("\n".join(chunks)))] if chunks else []


class _VisibleTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.hidden_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style", "noscript", "svg"}:
            self.hidden_depth += 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "br"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self.hidden_depth:
            self.hidden_depth -= 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden_depth:
            self.parts.append(data)


def extract_html(content: bytes) -> list[tuple[str, str]]:
    decoded = ""
    for encoding in ("utf-8", "big5", "cp950"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not decoded:
        decoded = content.decode("utf-8", errors="replace")
    parser = _VisibleTextParser()
    parser.feed(html.unescape(decoded))
    text = _clean_text("".join(parser.parts))
    return [("官方網頁全文", text)] if text else []


def extract_attachment(content: bytes, url: str, follow_pdf_links: bool = True) -> list[tuple[str, str]]:
    suffix = Path(url.split("?", 1)[0]).suffix.lower()
    if suffix == ".pdf" or content.startswith(b"%PDF"):
        if follow_pdf_links:
            pages = extract_pdf(content)
        else:
            with tempfile.NamedTemporaryFile(suffix=".pdf") as handle:
                handle.write(content)
                handle.flush()
                reader = PdfReader(handle.name)
                pages = [(f"Page {number}", _clean_text(page.extract_text() or ""))
                         for number, page in enumerate(reader.pages, start=1)]
                pages = [(name, text) for name, text in pages if text]
    elif suffix in {".odt", ".ods", ".docx"} or content.startswith(b"PK"):
        pages = extract_open_document(content)
    elif b"<html" in content[:2048].lower() or b"<!doctype html" in content[:2048].lower():
        pages = extract_html(content)
    else:
        raise ValueError(f"unsupported attachment type: {suffix or 'unknown'}")
    searchable = "\n".join(text for _, text in pages)
    if len(re.sub(r"\s+", "", searchable)) < MIN_SEARCHABLE_CHARS:
        raise ValueError("attachment contains no usable searchable text")
    return pages


def render_item(item: InventoryItem, pages: list[tuple[str, str]]) -> str:
    lines = [
        f"## {item.title}",
        f"來源網址：{item.url}",
        f"標籤：{item.tags}",
        f"上傳日期：{item.updated_date}",
        "",
    ]
    for page_name, text in pages:
        lines.extend((f"### {page_name}", text, ""))
    return "\n".join(lines).strip()


def build_markdown(items: list[InventoryItem], output_path: Path) -> None:
    sections = []
    failures = []
    for index, item in enumerate(items, start=1):
        try:
            pages = extract_attachment(download(item.url), item.url)
            sections.append(render_item(item, pages))
            print(f"[{index}/{len(items)}] OK {item.title}")
        except Exception as exc:  # keep the old complete file if any item is unreadable
            failures.append(f"{item.title}: {exc}")
            print(f"[{index}/{len(items)}] FAIL {item.title}: {exc}", file=sys.stderr)
    if failures:
        raise RuntimeError("refusing to write an incomplete corpus:\n" + "\n".join(failures))
    output_path.write_text("\n\n---\n\n".join(sections) + "\n", encoding="utf-8")
    print(f"wrote {len(sections)} readable documents to {output_path}")


PLACEHOLDER_RE = re.compile(
    r"（此法規僅取得檔案連結，未能抽出全文，請人工確認：(https?://[^）]+)）"
)


def repair_placeholders(path: Path) -> None:
    original = path.read_text(encoding="utf-8-sig")
    matches = list(PLACEHOLDER_RE.finditer(original))
    replacements = {}
    failures = []
    for index, match in enumerate(matches, start=1):
        url = match.group(1)
        try:
            pages = extract_attachment(download(url), url)
            replacement = f"來源網址：{url}\n\n" + "\n\n".join(
                f"### {page_name}\n{text}" for page_name, text in pages
            )
            replacements[match.group(0)] = replacement
            print(f"[{index}/{len(matches)}] repaired {url}")
        except Exception as exc:
            failures.append(f"{url}: {exc}")
            print(f"[{index}/{len(matches)}] FAIL {url}: {exc}", file=sys.stderr)
    if failures:
        raise RuntimeError("refusing to leave placeholder repairs incomplete:\n" + "\n".join(failures))
    repaired = original
    for old, new in replacements.items():
        repaired = repaired.replace(old, new)
    if any(marker in repaired for marker in FAILED_TEXT_MARKERS):
        raise RuntimeError(f"{path} still contains an unreadable-document marker")
    path.write_text(repaired, encoding="utf-8")
    print(f"repaired {len(matches)} placeholder documents in {path}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--target", action="append", choices=[*TARGETS, "ge-extra"])
    parser.add_argument("--repair-osa", action="store_true")
    args = parser.parse_args()
    targets = args.target or ["hr", "oga", "ge-extra"]
    for target in targets:
        if target == "ge-extra":
            build_markdown(load_missing_ge_items(), DATA / "ge_regulations_extra.md")
        else:
            dept_label, output_path = TARGETS[target]
            build_markdown(load_admin_items(dept_label), output_path)
    if args.repair_osa:
        repair_placeholders(DATA / "osa_regulations.md")


if __name__ == "__main__":
    main()
