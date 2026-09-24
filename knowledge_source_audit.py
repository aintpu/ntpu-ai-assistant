"""Static knowledge-source coverage checks for every supported NTPU office."""

from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "crawler_data"
ADMIN_XLSX = DATA / "北大行政單位法規彙整.xlsx"
ACADEMIC_XLSX = DATA / "北大學術單位法規彙整.xlsx"
FAILED_TEXT_MARKERS = ("未能抽出全文", "無法擷取內文", "擷取失敗")

OFFICE_SPECS = {
    "ope": {
        "inventory": ("admin", "體育室"),
        "sources": (DATA / "ALL_files_2.md",),
    },
    "ge": {
        "inventory": ("academic", "通識教育中心"),
        "sources": (DATA / "cge_content.md", DATA / "ge_regulations_extra.md"),
    },
    "lc": {
        "inventory": ("academic", "語言中心"),
        "sources": (DATA / "lc_content.md",),
    },
    "oaa": {
        "inventory": ("admin", "教務處"),
        "sources": (DATA / "oaa_regulations.md",),
    },
    "osa": {
        "inventory": ("admin", "學生事務處"),
        "sources": (DATA / "osa_regulations.md",),
    },
    "hr": {
        "inventory": ("admin", "人事室"),
        "sources": (DATA / "hr_regulations.md",),
    },
    "oga": {
        "inventory": ("admin", "總務處"),
        "sources": (DATA / "oga_regulations.md",),
    },
}


@dataclass(frozen=True)
class AuditResult:
    office: str
    inventory_rows: int
    readable_rows: int
    missing_titles: tuple[str, ...]
    unreadable_titles: tuple[str, ...]

    @property
    def passed(self) -> bool:
        return not self.missing_titles and not self.unreadable_titles


def normalize_title(value: str) -> str:
    value = str(value or "").strip()
    value = re.sub(r"^📄\s*", "", value)
    value = re.sub(r"\.(pdf|docx?|odt|ods)$", "", value, flags=re.I)
    return re.sub(r"\s+", "", value)


def _inventory_titles(kind: str, label: str) -> list[str]:
    path = ADMIN_XLSX if kind == "admin" else ACADEMIC_XLSX
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]] if kind == "admin" else workbook[label]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(rows)]
        titles = []
        for row in rows:
            record = dict(zip(headers, row))
            if kind == "admin" and str(record.get("處室") or "").strip() != label:
                continue
            title = str(record.get("法規名稱" if kind == "admin" else "title") or "").strip()
            if title:
                titles.append(title)
        return titles
    finally:
        workbook.close()


def _sections(paths: tuple[Path, ...]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    for path in paths:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8-sig")
        matches = list(re.finditer(r"^##\s+(.+?)\s*$", text, re.M))
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            title = normalize_title(match.group(1))
            sections.setdefault(title, []).append(text[match.end():end].strip())
    return sections


def _explicitly_unreadable_regulation_titles(paths: tuple[Path, ...]) -> set[str]:
    titles = set()
    for path in paths:
        if not path.exists() or not ("regulations" in path.name or path.name == "ALL_files_2.md"):
            continue
        text = path.read_text(encoding="utf-8-sig")
        matches = list(re.finditer(r"^##\s+(.+?)\s*$", text, re.M))
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            body = text[match.end():end]
            if any(marker in body for marker in FAILED_TEXT_MARKERS):
                titles.add(re.sub(r"^📄\s*", "", match.group(1)).strip())
    return titles


def _readable(body: str) -> bool:
    if any(marker in body for marker in FAILED_TEXT_MARKERS):
        return False
    content = re.sub(r"(?m)^(來源網址|標籤|上傳日期|文件類型|原始格式)：.*$", "", body)
    content = re.sub(r"(?m)^###\s+.*$", "", content)
    # A title or redirect label can easily exceed 20 characters.  Requiring a
    # larger body catches link-wrapper PDFs that have no answerable content.
    return len(re.sub(r"\s+", "", content)) >= 100


def audit_office(office: str) -> AuditResult:
    spec = OFFICE_SPECS[office]
    titles = _inventory_titles(*spec["inventory"])
    sections = _sections(spec["sources"])
    missing = []
    unreadable = []
    readable_count = 0
    for title in titles:
        bodies = sections.get(normalize_title(title), [])
        if not bodies:
            missing.append(title)
        elif not any(_readable(body) for body in bodies):
            unreadable.append(title)
        else:
            readable_count += 1
    unreadable.extend(sorted(
        _explicitly_unreadable_regulation_titles(spec["sources"]) - set(unreadable)
    ))
    return AuditResult(
        office=office,
        inventory_rows=len(titles),
        readable_rows=readable_count,
        missing_titles=tuple(missing),
        unreadable_titles=tuple(unreadable),
    )


def audit_all() -> list[AuditResult]:
    return [audit_office(office) for office in OFFICE_SPECS]


if __name__ == "__main__":
    results = audit_all()
    print(json.dumps([asdict(result) | {"passed": result.passed} for result in results],
                     ensure_ascii=False, indent=2))
    raise SystemExit(0 if all(result.passed for result in results) else 1)
